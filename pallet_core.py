# -*- coding: utf-8 -*-
"""
Created on Thu Feb 20 05:26:10 2025

@author: 陳皓均

棧板明細估算系統核心功能
"""
import json
import os
import re
import sys
from dataclasses import dataclass
from datetime import date

import pandas as pd
from openpyxl import load_workbook

RULE_COLUMNS = ["包材", "分裝", "包裝", "單一包裝"]
ITEM_COLUMNS = ["物料", "批次", "件數"]
OUTPUT_COLUMNS = ["品號", "批號", "數量", "包材"]
KIND_COLUMN = "類型"  # 只用於介面標色
FULL, REMAINDER, UNMATCHED = "滿分裝", "餘數", "查無規則"
OUTPUT_SHEET = "工作表2"
CODE_LENGTH = 5  # 物料代碼取最後幾碼比對包材
MAX_PROBLEMS = 3  # 錯誤訊息最多列出幾處

DATE_FORMAT = "%Y%m%d"
DATE_SUFFIX = re.compile(r"_\d{8}(_\d+)?$")  # 檔名結尾的日期與流水號

# 預設帶入的檔名
DEFAULT_RULE = "規則檔.xlsx"
DEFAULT_ITEM = "讀取檔.xlsx"
DEFAULT_OUTPUT = "寫入檔.xlsx"

SETTINGS_FILE = "pallet_settings.json"  # 記錄上次選擇的檔案


# 顯示在介面上的錯誤
class PalletError(Exception):
    pass


# 估算結果
@dataclass
class PalletResult:
    rows: pd.DataFrame  # 棧板明細
    item_count: int  # 合併後的品項數
    unmatched: list  # 找不到包材規則的物料
    output_path: str

    # 明細筆數
    @property
    def row_count(self):
        return len(self.rows)


# 檔案位置

# 程式所在資料夾(exe也適用)
def app_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


# 取得程式旁邊的預設檔案，不存在回傳空字串
def default_path(name):
    path = os.path.join(app_dir(), name)
    return path if os.path.exists(path) else ""


# 上次的選擇

# 讀取上次的選擇，失敗回傳空字典
def load_settings():
    path = os.path.join(app_dir(), SETTINGS_FILE)
    try:
        with open(path, encoding="utf-8") as file:
            settings = json.load(file)
    except (OSError, ValueError):
        return {}
    return settings if isinstance(settings, dict) else {}


# 儲存這次的選擇
def save_settings(settings):
    path = os.path.join(app_dir(), SETTINGS_FILE)
    try:
        with open(path, "w", encoding="utf-8") as file:
            json.dump(settings, file, ensure_ascii=False, indent=2)
    except OSError:
        pass


# 取得開啟時要帶入的檔案
def initial_path(saved_path, default_name):
    if isinstance(saved_path, str) and os.path.exists(saved_path):
        return saved_path
    return default_path(default_name)


# 產生今天日期的新檔路徑，重複時加流水號
def dated_path(template_path):
    folder, name = os.path.split(template_path)
    stem, ext = os.path.splitext(name)
    stem = DATE_SUFFIX.sub("", stem)
    base = f"{stem}_{date.today().strftime(DATE_FORMAT)}"

    path = os.path.join(folder, f"{base}{ext}")
    number = 2
    while os.path.exists(path):
        path = os.path.join(folder, f"{base}_{number}{ext}")
        number += 1
    return path


# 決定要寫入的檔案路徑
def output_path_for(template_path, keep_name):
    return template_path if keep_name else dated_path(template_path)


# 讀取檔案

# 讀取 Excel 並確認必要欄位都存在
def read_excel_checked(path, columns, label, **kwargs):
    name = os.path.basename(path)
    if not path:
        raise PalletError(f"請選擇{label}的檔案")
    if not os.path.exists(path):
        raise PalletError(f"{name} 不存在，請重新選擇")
    try:
        df = pd.read_excel(path, engine="openpyxl", **kwargs)
    except PermissionError:
        raise PalletError(f"{name} 被占用，請先關閉檔案再試")
    except Exception:
        raise PalletError(f"{name} 讀取失敗，請檢查檔案格式")

    missing = [col for col in columns if col not in df.columns]
    if missing:
        raise PalletError(f"{name} 缺少欄位：{'、'.join(missing)}")
    return df


# 檢查資料

# 判斷是否為正整數
def is_positive_int(value):
    number = pd.to_numeric(value, errors="coerce")
    return pd.notna(number) and number > 0 and float(number).is_integer()


# 判斷儲存格是否空白
def is_blank(value):
    return pd.isna(value) or not str(value).strip()


# 產生數值欄位的錯誤說明
def number_problem(row, column, value):
    if is_blank(value):
        return f"第{row}列{column}空白"
    number = pd.to_numeric(value, errors="coerce")
    if pd.notna(number) and float(number).is_integer():
        value = int(number)
    return f"第{row}列{column}「{value}」不是正整數"


# 有問題的列組成錯誤訊息
def raise_problems(path, problems):
    if not problems:
        return
    text = "；".join(problems[:MAX_PROBLEMS])
    if len(problems) > MAX_PROBLEMS:
        text += f" 等 {len(problems)} 處"
    raise PalletError(f"{os.path.basename(path)} {text}")


# 讀取並檢查條件規則
def load_rules(path):
    df = read_excel_checked(path, RULE_COLUMNS, "條件規則",
                            dtype={"包材": str})
    df = df.dropna(how="all", subset=RULE_COLUMNS)

    problems = []
    for index, code, size in zip(df.index, df["包材"], df["分裝"]):
        row = index + 2
        if is_blank(code):
            problems.append(f"第{row}列包材空白")
        if not is_positive_int(size):
            problems.append(number_problem(row, "分裝", size))
    raise_problems(path, problems)

    df["分裝"] = pd.to_numeric(df["分裝"]).astype(int)
    return df


# 讀取並檢查需要整合的資料
def load_items(path):
    df = read_excel_checked(path, ITEM_COLUMNS, "需要整合",
                            dtype={"物料": str, "批次": str})
    df = df.dropna(how="all", subset=ITEM_COLUMNS)

    problems = []
    for index, item_code, qty in zip(df.index, df["物料"], df["件數"]):
        row = index + 2
        if is_blank(item_code):
            problems.append(f"第{row}列物料空白")
        if not is_positive_int(qty):
            problems.append(number_problem(row, "件數", qty))
    raise_problems(path, problems)

    df["件數"] = pd.to_numeric(df["件數"]).astype(int)
    return df


# 確認輸出檔可以開啟
def check_output(path):
    name = os.path.basename(path)
    if not path:
        raise PalletError("請選擇輸出結果的檔案")
    if not os.path.exists(path):
        raise PalletError(f"{name} 不存在，請重新選擇")
    try:
        load_workbook(path, read_only=True).close()
    except PermissionError:
        raise PalletError(f"{name} 被占用，請先關閉檔案再試")
    except Exception:
        raise PalletError(f"{name} 開啟失敗，請檢查檔案格式")


# 拆分明細

# 合併同物料、批次的件數
def merge_items(df):
    df["批次"] = df["批次"].fillna("")
    return df.groupby(["物料", "批次"], sort=False, as_index=False)["件數"].sum()


# 找出物料對應的包材規則，查無回傳None
def find_rule(rules_df, item_code):
    tail = item_code[-CODE_LENGTH:]
    matched = rules_df[rules_df["包材"].apply(lambda code: code in tail)]
    if matched.empty:
        return None
    return matched.iloc[0]


# 依分裝數量拆成多筆明細
def split_items(items_df, rules_df):
    result = []
    unmatched = []
    rows = items_df[ITEM_COLUMNS].itertuples(index=False)
    for item_code, batch, qty in rows:
        rule = find_rule(rules_df, item_code)
        if rule is None:
            result.append([item_code, batch, qty, "", UNMATCHED])
            if item_code not in unmatched:
                unmatched.append(item_code)
            continue

        quotient, remainder = divmod(int(qty), int(rule["分裝"]))

        full_row = [item_code, batch, rule["分裝"], rule["包裝"], FULL]
        result.extend([full_row] * quotient)
        if remainder > 1:
            result.append([item_code, batch, remainder, rule["包裝"],
                           REMAINDER])
        elif remainder == 1:
            result.append([item_code, batch, remainder, rule["單一包裝"],
                           REMAINDER])

    columns = OUTPUT_COLUMNS + [KIND_COLUMN]
    return pd.DataFrame(result, columns=columns), unmatched


# 寫入檔案

# 轉成 Excel 儲存格可接受的值
def cell_value(value):
    if pd.isna(value):
        return None
    return value.item() if hasattr(value, "item") else value


# 把明細寫入工作表2並存檔
def write_output(output_df, template_path, save_path):
    name = os.path.basename(save_path)
    check_output(template_path)
    try:
        workbook = load_workbook(template_path)
        if OUTPUT_SHEET in workbook.sheetnames:
            sheet = workbook[OUTPUT_SHEET]
        else:
            sheet = workbook.create_sheet(OUTPUT_SHEET)

        width = len(OUTPUT_COLUMNS)
        for row in sheet.iter_rows(min_row=2, max_col=width):
            for cell in row:
                cell.value = None

        for col, title in enumerate(OUTPUT_COLUMNS, start=1):
            sheet.cell(row=1, column=col, value=title)
        rows = output_df[OUTPUT_COLUMNS].itertuples(index=False)
        for r, values in enumerate(rows, start=2):
            for col, value in enumerate(values, start=1):
                sheet.cell(row=r, column=col, value=cell_value(value))

        workbook.save(save_path)
    except PermissionError:
        raise PalletError(f"{name} 被占用，請先關閉檔案再試")
    except Exception:
        raise PalletError(f"{name} 寫入失敗，請檢查檔案格式")


# 執行估算並寫檔
def build_pallet(rule_path, item_path, template_path, keep_name=False):
    rules_df = load_rules(rule_path)
    items_df = load_items(item_path)
    check_output(template_path)

    merged_df = merge_items(items_df)
    output_df, unmatched = split_items(merged_df, rules_df)
    save_path = output_path_for(template_path, keep_name)
    write_output(output_df, template_path, save_path)
    return PalletResult(output_df, len(merged_df), unmatched, save_path)
